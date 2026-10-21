"""AutoProctor admin-site scraper — pure I/O, no database. The nightly sync
(autoproctor_service.run_sync) drives it.

How the site works, as observed 2026-10-17:
  * Login is passwordless: enter the admin email, AutoProctor mails a 6-digit
    code from hello@autoproctor.co with the code in the SUBJECT
    ("935806 is your AutoProctor Login Code"), type it into six one-digit
    boxes, click Verify. We read that mail through MS Graph (Mail.Read).
  * The site's load balancer returns 403 to anything that announces itself as
    a bot or headless browser, so the browser runs with a regular Chrome UA.
  * Its JSON API needs a Bearer token the web app attaches itself, so rather
    than calling it directly we drive the pages and read the JSON responses
    the pages fetch (page.on("response")).
  * Latest attempts: /api/v3/test-attempts/?…&limit=100, newest first, only
    finished attempts; scrolling to the bottom loads the next 100.
  * Report page: Results / Proctoring / Auxiliary device tabs. "Download All
    responses" (Results) is a real .xlsx download, but it holds EVERY
    candidate's responses for that test, so we fetch it once per test per run
    and slice out each candidate. The two "Download Report" buttons just call
    window.print(); we used to reproduce them with page.pdf(), but loading and
    printing ~150 evidence photos crashed Chromium on a 1 GB worker, so the
    PDFs were dropped (2026-10-17) — HR opens the evidence in AutoProctor.
"""
import asyncio
import io
import logging
import re
from datetime import datetime, timezone, timedelta
from typing import Optional

import httpx

from app.config import settings

logger = logging.getLogger(__name__)

BASE = "https://www.autoproctor.co"
LOGIN_URL = f"{BASE}/account/login/role/test-admin?next=https%3A%2F%2Fwww.autoproctor.co%2Ftest-admin%2Flatest-attempts%2F"
LATEST_URL = f"{BASE}/test-admin/latest-attempts/"
CHROME_UA = (
    "Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) AppleWebKit/537.36 "
    "(KHTML, like Gecko) Chrome/129.0.0.0 Safari/537.36"
)
OTP_REGEX = re.compile(r"\b(\d{6})\b")

# Proctoring-tab counter labels -> our violation keys. Matched loosely
# (case-insensitive, any suffix) because the page and its print layout word
# them slightly differently ("TAB SWITCHES" vs "TAB SWITCHED").
VIOLATION_LABELS = {
    "tab_switched": r"tab\s+switch\w*",
    "no_face": r"no\s+face\s+detected",
    "multiple_faces": r"multiple\s+faces",
    "noise": r"noise\s+detected",
    "exited_fullscreen": r"exited\s+full\s*screen",
    "multiple_monitors": r"multiple\s+monitors",
}


class AutoproctorError(Exception):
    pass


def missing_config() -> list[str]:
    """Settings the sync can't run without. An empty login email made every
    login fail with "Email cannot be empty" (2026-10-18); an empty OTP sender
    would let ANY mail with a 6-digit subject pass as the login code."""
    required = {
        "AUTOPROCTOR_LOGIN_EMAIL": settings.AUTOPROCTOR_LOGIN_EMAIL,
        "AUTOPROCTOR_OTP_SENDER": settings.AUTOPROCTOR_OTP_SENDER,
    }
    out = [k for k, v in required.items() if not (v or "").strip()]
    if not settings.ms_graph_configured:
        out.append("MS_GRAPH_TENANT_ID / MS_GRAPH_CLIENT_ID / MS_GRAPH_CLIENT_SECRET")
    return out


def report_url(attempt_label: str) -> str:
    return f"{BASE}/test-admin/test-attempts/{attempt_label}/report/?t=p"


def _parse_dt(value: Optional[str]) -> Optional[datetime]:
    if not value:
        return None
    try:
        dt = datetime.fromisoformat(value.replace("Z", "+00:00"))
    except ValueError:
        return None
    return dt if dt.tzinfo else dt.replace(tzinfo=timezone.utc)


def _hundredths(value) -> Optional[float]:
    return None if value is None else round(value / 100, 2)


# ── Login code from the mailbox ──────────────────────────────────────────────

async def fetch_login_otp(requested_at: datetime, timeout_s: int = 150, poll_s: int = 4) -> str:
    """Wait for an AutoProctor login code that arrived after requested_at.

    requested_at must be captured BEFORE the "Sign In With Email" click, so a
    code from an earlier login can never be picked up."""
    from app.services.ms_graph_service import GRAPH_BASE, _get_app_token

    if not settings.ms_graph_configured:
        raise AutoproctorError("MS Graph is not configured — can't read the login code")

    cutoff = (requested_at - timedelta(seconds=10)).strftime("%Y-%m-%dT%H:%M:%SZ")
    url = f"{GRAPH_BASE}/users/{settings.AUTOPROCTOR_LOGIN_EMAIL}/messages"
    params = {
        "$filter": f"receivedDateTime ge {cutoff}",
        "$orderby": "receivedDateTime desc",
        "$top": "10",
        "$select": "from,subject,bodyPreview,receivedDateTime",
    }
    sender = (settings.AUTOPROCTOR_OTP_SENDER or "").strip().lower()
    if not sender:
        raise AutoproctorError("AUTOPROCTOR_OTP_SENDER isn't set — refusing to guess which email holds the code")
    deadline = asyncio.get_event_loop().time() + timeout_s
    async with httpx.AsyncClient(timeout=20) as client:
        while asyncio.get_event_loop().time() < deadline:
            token = await _get_app_token()
            resp = await client.get(url, params=params, headers={"Authorization": f"Bearer {token}"})
            resp.raise_for_status()
            for msg in resp.json().get("value", []):
                frm = ((msg.get("from") or {}).get("emailAddress") or {}).get("address", "").lower()
                if sender not in frm:
                    continue
                for text in (msg.get("subject"), msg.get("bodyPreview")):
                    m = OTP_REGEX.search(text or "")
                    if m:
                        logger.info(f"AutoProctor login code received at {msg.get('receivedDateTime')} "
                                    f"(****{m.group(1)[-2:]})")
                        return m.group(1)
            await asyncio.sleep(poll_s)
    raise AutoproctorError(f"No AutoProctor login code reached {settings.AUTOPROCTOR_LOGIN_EMAIL} within {timeout_s}s")


# ── Responses workbook ───────────────────────────────────────────────────────

def slice_responses(xlsx: bytes, email: str) -> tuple[Optional[bytes], list[dict]]:
    """Cut one candidate out of a test's "Download All responses" workbook.

    Returns (a workbook with just their rows + the Questions sheet, a
    question-by-question list for the portal). (None, []) when the candidate
    isn't in the file — e.g. AutoProctor hadn't flushed it yet."""
    import openpyxl

    wb = openpyxl.load_workbook(io.BytesIO(xlsx), read_only=True, data_only=True)
    sheets = {ws.title: list(ws.iter_rows(values_only=True)) for ws in wb.worksheets}
    email = email.strip().lower()

    def candidate_row(rows):
        if not rows:
            return None, None
        header = list(rows[0])
        for r in rows[1:]:
            if r and isinstance(r[0], str) and r[0].strip().lower() == email:
                return header, list(r)
        return header, None

    points_header, points_row = candidate_row(sheets.get("Points", []))
    resp_header, resp_row = candidate_row(sheets.get("Responses", []))
    if points_row is None and resp_row is None:
        return None, []

    questions = {}
    for r in (sheets.get("Questions") or [])[1:]:
        if r and r[0] is not None:
            questions[str(r[0]).strip()] = {"question": r[1], "max_points": r[2]}

    def by_number(header, row):
        out = {}
        if not header or not row:
            return out
        for h, v in zip(header, row):
            m = re.match(r"^(\d+)\.\s", str(h or ""))
            if m:
                out[m.group(1)] = v
        return out

    pts = by_number(points_header, points_row)
    ans = by_number(resp_header, resp_row)
    numbers = sorted(set(questions) | set(pts) | set(ans), key=lambda n: int(n) if n.isdigit() else 0)
    responses = [
        {
            "number": int(n) if n.isdigit() else n,
            "question": (questions.get(n) or {}).get("question"),
            "answer": None if ans.get(n) is None else str(ans.get(n)),
            "points": pts.get(n),
            "max_points": (questions.get(n) or {}).get("max_points"),
        }
        for n in numbers
    ]

    out = openpyxl.Workbook()
    out.remove(out.active)
    for title, header, row in (("Points", points_header, points_row), ("Responses", resp_header, resp_row)):
        if header:
            ws = out.create_sheet(title)
            ws.append(header)
            if row:
                ws.append(row)
    if sheets.get("Questions"):
        ws = out.create_sheet("Questions")
        for r in sheets["Questions"]:
            ws.append(list(r))
    buf = io.BytesIO()
    out.save(buf)
    return buf.getvalue(), responses


# ── Browser session ──────────────────────────────────────────────────────────

# The logged-in browser session (cookies) is kept in Redis between runs, so
# a run only requests a login code when the saved session has expired.
# Requesting a code on every run got us cut off: AutoProctor stopped sending
# codes after seven in a few hours (2026-10-17), failing the next run's login.
SESSION_KEY = "autoproctor:session"
SESSION_TTL_SECONDS = 30 * 24 * 3600


async def _load_session() -> Optional[dict]:
    try:
        import json
        from app.services.auth_service import _redis_client
        raw = await _redis_client().get(SESSION_KEY)
        return json.loads(raw) if raw else None
    except Exception as exc:
        logger.warning(f"AutoProctor: couldn't load the saved session ({exc}); will log in")
        return None


async def _save_session(state: dict) -> None:
    try:
        import json
        from app.services.auth_service import _redis_client
        await _redis_client().set(SESSION_KEY, json.dumps(state), ex=SESSION_TTL_SECONDS)
    except Exception as exc:
        logger.warning(f"AutoProctor: couldn't save the session ({exc}); next run will log in again")


async def _clear_session() -> None:
    try:
        from app.services.auth_service import _redis_client
        await _redis_client().delete(SESSION_KEY)
    except Exception:
        pass


# Words that mark AutoProctor's own error/notice text on the login form.
_NOTICE_RE = re.compile(r"(too many|try again|limit|wait|invalid|error|blocked|unable|later|cannot|empty|required|not allowed)", re.I)


class AutoproctorSession:
    """`async with AutoproctorSession() as ap:` launches Chromium and gets a
    logged-in session: the saved one if it's still valid, otherwise a fresh
    email-code login."""

    def __init__(self):
        self._pw = None
        self._browser = None
        self.ctx = None

    async def __aenter__(self):
        from playwright.async_api import async_playwright

        missing = missing_config()
        if missing:
            raise AutoproctorError(f"AutoProctor sync isn't configured on this worker — set: {', '.join(missing)}")

        self._pw = await async_playwright().start()
        # --no-sandbox: the container runs as root, where Chromium's sandbox
        # refuses to start. --disable-dev-shm-usage: /dev/shm is tiny in
        # containers and big report pages crash the renderer without it.
        self._browser = await self._pw.chromium.launch(
            headless=True, args=["--no-sandbox", "--disable-dev-shm-usage"],
        )
        saved = await _load_session()
        self.ctx = await self._browser.new_context(
            user_agent=CHROME_UA, viewport={"width": 1440, "height": 900}, accept_downloads=True,
            storage_state=saved,
        )
        self.ctx.set_default_timeout(60_000)
        if saved and await self._session_is_valid():
            logger.info("AutoProctor: reusing saved session, no login code needed")
        else:
            await self._login()
        await _save_session(await self.ctx.storage_state())
        return self

    async def __aexit__(self, *exc):
        try:
            if self.ctx:
                # Cookies may have been refreshed during the run.
                await _save_session(await self.ctx.storage_state())
        except Exception:
            pass
        try:
            if self._browser:
                await self._browser.close()
        finally:
            if self._pw:
                await self._pw.stop()

    async def _session_is_valid(self) -> bool:
        """An expired session gets redirected to the login page."""
        page = await self.ctx.new_page()
        try:
            await page.goto(LATEST_URL, wait_until="domcontentloaded")
            await page.wait_for_timeout(1500)
            if "/account/login" in page.url:
                await _clear_session()
                logger.info("AutoProctor: saved session expired, logging in again")
                return False
            return True
        except Exception:
            return False
        finally:
            await page.close()

    async def _fill_email(self, page) -> None:
        """Type the login email and confirm the field really holds it before
        submitting, so a bad fill fails here with a clear message instead of
        as AutoProctor's "Email cannot be empty" toast."""
        email = settings.AUTOPROCTOR_LOGIN_EMAIL.strip()
        field = page.locator("#email-desktop")
        await field.wait_for(state="visible")
        for _ in range(3):
            await field.fill(email)
            if (await field.input_value()).strip().lower() == email.lower():
                return
            await page.wait_for_timeout(1000)
        raise AutoproctorError("Couldn't type the login email into AutoProctor's sign-in form")

    async def _wait_for_code_screen(self, page) -> Optional[str]:
        """After "Sign In With Email": None once the code screen is up, or the
        error AutoProctor showed instead. Raises if neither appears in 60s."""
        verify = page.get_by_role("button", name="Verify", exact=True)
        deadline = asyncio.get_event_loop().time() + 60
        while asyncio.get_event_loop().time() < deadline:
            if await verify.count() and await verify.first.is_visible():
                return None
            notice = await self._page_notice(page)
            if notice:
                return notice
            await page.wait_for_timeout(1000)
        raise AutoproctorError("AutoProctor didn't show the login-code screen within 60s")

    async def _login(self):
        page = await self.ctx.new_page()
        try:
            await page.goto(LOGIN_URL, wait_until="domcontentloaded")
            # Let the page's scripts (form hydration, cookie banner) settle.
            try:
                await page.wait_for_load_state("networkidle", timeout=15_000)
            except Exception:
                pass  # analytics can keep the network busy; the checks below cover it
            try:
                await page.get_by_role("button", name="OK, I accept").click(timeout=4000)
                await page.wait_for_timeout(1000)  # accepting can re-render the page
            except Exception:
                pass  # cookie banner only shows on a fresh context, and not always
            await self._fill_email(page)
            requested_at = datetime.now(timezone.utc)
            await page.click("#sign-in-button-desktop")
            notice = await self._wait_for_code_screen(page)
            if notice is not None:
                raise AutoproctorError(f'AutoProctor refused the login request: "{notice}"')

            code = await fetch_login_otp(requested_at)
            boxes = [
                el for el in await page.query_selector_all("input")
                if await el.is_visible() and (await el.get_attribute("type")) not in ("hidden", "email", "checkbox")
            ]
            if len(boxes) >= 6:
                for box, digit in zip(boxes, code):
                    await box.fill(digit)
            elif boxes:
                await boxes[0].fill(code)
            else:
                raise AutoproctorError("Login code boxes not found on the verify page")
            await page.get_by_role("button", name="Verify", exact=True).click()
            await page.wait_for_url("**/test-admin/**", timeout=60_000)
            logger.info("AutoProctor login ok")
        finally:
            await page.close()

    @staticmethod
    async def _page_notice(page) -> Optional[str]:
        """The error/notice line the login form is showing, if any."""
        try:
            lines = [ln.strip() for ln in (await page.inner_text("body")).splitlines()]
            hits = [ln for ln in lines if 3 < len(ln) < 200 and _NOTICE_RE.search(ln)]
            return " | ".join(dict.fromkeys(hits))[:300] or None
        except Exception:
            return None

    # ── Latest attempts ──

    async def list_attempts(self, since: datetime, max_pages: int = 30) -> list[dict]:
        """Every finished attempt that STARTED on/after `since`, newest first,
        as AutoProctor's list rows. Scrolls for more until the oldest loaded
        row predates `since` or the list runs out."""
        pages: list[dict] = []
        got = asyncio.Event()

        async def on_response(resp):
            if "/api/v3/test-attempts/?" in resp.url and resp.request.method == "GET":
                try:
                    pages.append(await resp.json())
                    got.set()
                except Exception:
                    pass

        page = await self.ctx.new_page()
        page.on("response", on_response)
        try:
            await page.goto(LATEST_URL, wait_until="domcontentloaded")
            for _ in range(max_pages):
                try:
                    await asyncio.wait_for(got.wait(), timeout=30)
                except asyncio.TimeoutError:
                    break
                got.clear()
                last = pages[-1]
                rows = last.get("data") or []
                oldest = _parse_dt(rows[-1].get("test_client_start_time")) if rows else None
                if not last.get("has_more") or (oldest and oldest < since):
                    break
                await page.evaluate("window.scrollTo(0, document.body.scrollHeight)")
        finally:
            await page.close()

        seen, out = set(), []
        for p in pages:
            for row in p.get("data") or []:
                label = row.get("label")
                started = _parse_dt(row.get("test_client_start_time"))
                if not label or label in seen or not row.get("test_finish_time"):
                    continue
                if started and started < since:
                    continue
                seen.add(label)
                out.append(row)
        return out

    # ── One attempt ──

    async def _open_report(self, attempt_label: str):
        captured: dict[str, dict] = {}

        async def on_response(resp):
            url = resp.url
            if f"/test-attempts/{attempt_label}/" not in url:
                return
            try:
                data = await resp.json()
            except Exception:
                return
            if "/api/v3/" in url and "/report" not in url:
                captured["v3"] = data
            elif re.search(r"/api/v2/test-attempts/[^/]+/?$", url.split("?")[0]):
                captured["v2"] = data

        page = await self.ctx.new_page()
        page.on("response", on_response)
        # The report page streams evidence media and never goes network-idle.
        await page.goto(report_url(attempt_label), wait_until="domcontentloaded")
        await page.get_by_role("tab", name="Proctoring").wait_for(timeout=60_000)
        for _ in range(30):
            if "v3" in captured:
                break
            await page.wait_for_timeout(500)
        await page.wait_for_timeout(3000)
        return page, captured

    @staticmethod
    async def _read_violations(page) -> dict:
        text = await page.inner_text("body")
        out = {}
        for key, pattern in VIOLATION_LABELS.items():
            m = re.search(pattern + r"\s*\n?\s*(\d+)", text, flags=re.I)
            out[key] = int(m.group(1)) if m else None
        return out

    async def scrape_attempt(self, attempt_label: str) -> dict:
        """Top-level data and the proctoring violation counters. The report
        PDFs are deliberately not generated: printing them crashed Chromium
        out of memory on the 1 GB worker, and the photo evidence is one click
        away in AutoProctor itself."""
        page, captured = await self._open_report(attempt_label)
        try:
            v3 = captured.get("v3") or {}
            v3 = v3.get("data", v3)  # the detail endpoint wraps its payload in {"data": …}
            if not v3.get("test_taker"):
                raise AutoproctorError(f"Attempt {attempt_label}: report data never loaded")
            await page.get_by_role("tab", name="Proctoring").click()
            # The counters render after the tab's own data arrives — wait for
            # them rather than a fixed sleep, or a slow load reads them as blank.
            try:
                await page.wait_for_function(
                    "() => /tab\\s+switch\\w*\\s+\\d+/i.test(document.body.innerText)", timeout=30_000,
                )
            except Exception:
                logger.warning(f"Attempt {attempt_label}: violation counters didn't render")
            violations = await self._read_violations(page)
        finally:
            await page.close()

        v2 = (captured.get("v2") or {}).get("test_attempt") or {}
        test = v3.get("test") or {}
        taker = v3.get("test_taker") or {}
        name = " ".join(x for x in (taker.get("first_name"), taker.get("last_name")) if x).strip()
        return {
            "attempt_label": attempt_label,
            "test_label": test.get("label"),
            "test_title": test.get("title") or "",
            "candidate_email": (taker.get("email") or "").strip().lower(),
            "candidate_name": name or None,
            "started_at": _parse_dt(v3.get("test_client_start_time")),
            "finished_at": _parse_dt(v3.get("test_finish_time")),
            "duration_seconds": v3.get("duration_seconds"),
            "device_type": v3.get("device_type"),
            "trust_score_pct": v3.get("trust_score_pct"),
            "percentage_scored": _hundredths(v3.get("percentage_scored_times_hundred")),
            "points_scored": _hundredths(v3.get("points_scored_times_hundred")),
            "points_available": _hundredths(test.get("soc_points_available_times_hundred")),
            "graded": bool(v3.get("graded")),
            "violations": violations,
            "aux_device_verified": (v2.get("misc_data") or {}).get("auxiliary_device_verified"),
            # test_settings is a large static blob of plan/credit config — not worth keeping per row.
            "raw": {k: v for k, v in v3.items() if k != "test_settings"},
        }

    async def download_test_responses(self, attempt_label: str) -> bytes:
        """The test-wide "Download All responses" workbook, reached through
        any one attempt's Results tab."""
        page, _ = await self._open_report(attempt_label)
        try:
            await page.get_by_role("tab", name="Results").click()
            await page.wait_for_timeout(3000)
            async with page.expect_download(timeout=120_000) as info:
                await page.get_by_role("button", name="Download All responses").click()
            download = await info.value
            path = await download.path()
            with open(path, "rb") as fh:
                return fh.read()
        finally:
            await page.close()
